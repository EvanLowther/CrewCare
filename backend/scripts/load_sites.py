"""Load the ~160 demo sites (US power plants) from EIA-860M into `sites`.

    python scripts/load_sites.py                 # dry run: print the selection, write nothing
    python scripts/load_sites.py --write         # upsert into Tiger
    python scripts/load_sites.py --write --prune # ...and delete sites that aren't in the selection
    python scripts/load_sites.py path/to/file.xlsx [--write]

Get the latest monthly file from https://www.eia.gov/electricity/data/eia860m/ and save it in
data/raw/ (git-ignored). With no path, the newest data/raw/*generator*.xlsx is used.

EIA-860M lists generators, not plants: rows are grouped by Plant ID, capacity is summed, and the
largest generator's technology becomes the plant type. Safe to re-run (upsert on id).

The selection is the hand-picked Florida sites (SELECTED_SITES) plus NATIONAL_SITES more spread over every other
state: one per state, the rest by each state's share of large plants (allocate()), largest first and at least
MIN_SPACING_MI apart within a state (select_national()).
"""

from __future__ import annotations

import argparse
import math
import sys
from collections import Counter
from collections.abc import Collection, Iterable, Sequence
from dataclasses import dataclass
from functools import cache
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = REPO_ROOT / "data" / "raw"

# West of the Apalachicola River is Central time. Every selected site east of the river sits
# east of -85.0 and every one west of it sits west of -85.0, so a longitude cut is enough here.
# (timezonefinder puts Jim Woodruff Dam, on the river itself, in Central time; the plant keeps Eastern.)
CENTRAL_TIME_WEST_OF_LON = -85.0

# Sites outside Florida. Sized for Open-Meteo's free tier (10,000 calls/day): 100 more sites x 2 years of
# hourly history is ~5,200 calls, and 160 sites polled every 30 min at <= 10 variables is ~7,700 calls/day.
NATIONAL_SITES = 100
LARGE_PLANT_MW = 50.0  # states are weighted by how many plants this size they have; sites are picked from them
MIN_SPACING_MI = 50.0  # two sites in one state at least this far apart, so they see different weather
EARTH_RADIUS_MI = 3958.8

# Plant ID -> why it's in the demo. Picked from the August 2026 file for a mix of size, technology
# and hazards; keep near-duplicate sites out (e.g. Lauderdale shares a location with Dania Beach 7).
SELECTED_SITES: dict[int, str] = {
    # Largest plants: where most crews and equipment are
    56407: "largest plant in Florida",
    628: "coal + gas, Gulf coast storm surge",
    621: "nuclear, next to Biscayne Bay and the Everglades",
    6042: "large gas plant, Tampa Bay area",
    612: "on the Caloosahatchee River",
    6043: "large gas plant next to the Everglades",
    620: "on Lake Monroe / St. Johns River",
    7873: "Tampa Bay waterfront",
    7302: "large gas plant, central Florida",
    6045: "nuclear, barrier island on the Atlantic",
    136: "coal, on the St. Johns River",
    645: "Tampa Bay waterfront",
    60345: "north of Lake Okeechobee",
    634: "Tampa Bay waterfront",
    7242: "large gas plant, central Florida",
    65978: "Fort Lauderdale, dense urban site",
    667: "Jacksonville, St. Johns River mouth",
    617: "Port Everglades, Atlantic coast",
    619: "Riviera Beach, Atlantic coast",
    609: "Cape Canaveral, Atlantic coast",
    564: "coal + gas, Orlando",
    8049: "peaking plant, central Florida",
    8048: "Gulf coast, Pasco County",
    55286: "Brevard County, near the Indian River Lagoon",
    7380: "near the Peace River",
    7846: "Jacksonville, inland",
    6046: "peaking plant near the St. Johns River",
    55415: "near the Peace River",
    7238: "central Florida, Kissimmee",
    688: "Tallahassee",
    676: "Lakeland",
    663: "Gainesville",
    55318: "Indian River Lagoon",
    55412: "central Florida, Polk County",
    55414: "Pasco County, inland",
    56400: "Fort Pierce, Atlantic coast",
    56799: "Jacksonville, south",
    # Panhandle (Central time)
    641: "panhandle, Pensacola (Central time)",
    643: "panhandle, Panama City bay (Central time)",
    55242: "panhandle, Santa Rosa County (Central time)",
    65433: "panhandle solar on the Blackwater River (Central time)",
    65888: "panhandle solar, Okaloosa County (Central time)",
    65432: "panhandle solar near the Apalachicola basin (Central time)",
    # Rivers (flood warnings; USGS gauges get attached later)
    638: "on the Suwannee River",
    690: "Jim Woodruff Dam on the Apalachicola River",
    65871: "solar on the Caloosahatchee River",
    10611: "Palatka, on the St. Johns River",
    62750: "solar on the Little Manatee River",
    63518: "solar where the Suwannee and Withlacoochee meet",
    # South Florida / Everglades solar (heat, wildlife, swamp access roads)
    65423: "Everglades Solar Energy Center, Miami-Dade",
    65038: "Ghost Orchid solar, Big Cypress edge",
    65037: "Sawgrass solar, Big Cypress edge",
    68823: "Goldenrod solar, Collier County",
    62631: "Blue Heron solar, Hendry County",
    59993: "Babcock Ranch solar, Charlotte County",
    63110: "Sabal Palm solar, western Palm Beach County",
    62922: "solar on the Lake Okeechobee shore",
    # Everglades Agricultural Area and the Keys
    54627: "sugar-cane biomass, Everglades Agricultural Area",
    6584: "Key West: hurricanes and storm surge",
    10202: "biomass mill on the Amelia River, Fernandina Beach",
}


@dataclass(frozen=True)
class Plant:
    id: int
    name: str
    utility: str
    county: str
    state: str
    plant_type: str
    capacity_mw: float
    latitude: float
    longitude: float

    @property
    def timezone(self) -> str:
        return timezone_for(self.state, self.latitude, self.longitude)


@cache
def _timezone_finder():
    from timezonefinder import TimezoneFinder  # only needed outside Florida

    return TimezoneFinder()


def timezone_for(state: str, latitude: float, longitude: float) -> str:
    """IANA timezone for a site: Florida by the Apalachicola cut, everywhere else by timezone boundaries."""
    if state == "FL":
        return "America/Chicago" if longitude < CENTRAL_TIME_WEST_OF_LON else "America/New_York"
    tz = _timezone_finder().timezone_at(lng=longitude, lat=latitude)
    if tz is None or tz.startswith("Etc/"):  # offshore point outside every land zone
        raise ValueError(f"No timezone for {state} site at {latitude}, {longitude}")
    return tz


def miles(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    """Great-circle distance in miles."""
    p1, p2 = math.radians(lat1), math.radians(lat2)
    a = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(lon2 - lon1) / 2) ** 2
    return 2 * EARTH_RADIUS_MI * math.asin(math.sqrt(a))


def find_header_row(rows: Iterable[Sequence[object]]) -> int:
    """Index of the first row containing a 'Plant ID' cell (the sheet has title rows above it)."""
    for i, row in enumerate(rows):
        if any(isinstance(c, str) and c.strip() == "Plant ID" for c in row):
            return i
    raise ValueError("No 'Plant ID' header row found. Is this an EIA-860M generator file?")


def _num(value: object) -> float | None:
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None


def plants_from_generators(
    header: Sequence[object], rows: Iterable[Sequence[object]], states: Collection[str] | None = None
) -> dict[int, Plant]:
    """Group generator rows into plants, for the given state codes or (None) every state."""
    col = {str(name).strip(): i for i, name in enumerate(header) if name is not None}
    needed = ["Plant ID", "Plant Name", "Entity Name", "Plant State", "County", "Technology",
              "Nameplate Capacity (MW)", "Latitude", "Longitude"]
    missing = [n for n in needed if n not in col]
    if missing:
        raise ValueError(f"Missing columns in EIA file: {missing}")

    totals: dict[int, float] = {}
    largest: dict[int, tuple[float, Sequence[object]]] = {}
    for row in rows:
        if not row[col["Plant State"]] or (states is not None and row[col["Plant State"]] not in states):
            continue
        plant_id = _num(row[col["Plant ID"]])
        if plant_id is None:
            continue
        pid = int(plant_id)
        mw = _num(row[col["Nameplate Capacity (MW)"]]) or 0.0
        totals[pid] = totals.get(pid, 0.0) + mw
        if pid not in largest or mw > largest[pid][0]:
            largest[pid] = (mw, row)

    plants = {}
    for pid, (_, row) in largest.items():
        lat, lon = _num(row[col["Latitude"]]), _num(row[col["Longitude"]])
        if lat is None or lon is None:
            continue
        plants[pid] = Plant(
            id=pid,
            name=str(row[col["Plant Name"]]).strip(),
            utility=str(row[col["Entity Name"]]).strip(),
            county=str(row[col["County"]]).strip(),
            state=str(row[col["Plant State"]]).strip(),
            plant_type=str(row[col["Technology"]]).strip(),
            capacity_mw=round(totals[pid], 1),
            latitude=lat,
            longitude=lon,
        )
    return plants


def read_plants(path: Path) -> dict[int, Plant]:
    from openpyxl import load_workbook  # only needed when reading the real file

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = wb["Operating"].iter_rows(values_only=True)
    head = []
    for row in rows:  # consume title rows up to and including the header
        head.append(row)
        try:
            find_header_row([row])
            break
        except ValueError:
            continue
    return plants_from_generators(head[-1], rows)


def select(plants: dict[int, Plant]) -> tuple[list[Plant], list[int]]:
    """The selected Florida plants (in SELECTED_SITES order) and any selected ids missing from the file."""
    found = [plants[pid] for pid in SELECTED_SITES if pid in plants]
    missing = [pid for pid in SELECTED_SITES if pid not in plants]
    return found, missing


def allocate(counts: dict[str, int], total: int) -> dict[str, int]:
    """Sites per state: one each, the rest in proportion to `counts` (largest remainder), never more than a
    state has. Ties go to the larger state, then alphabetically, so the result is stable."""
    if total < len(counts):
        raise ValueError(f"{total} sites can't cover {len(counts)} states")
    alloc = dict.fromkeys(counts, 1)
    # Repeats only when a small state hits its cap and its share has to go to the others.
    while (left := total - sum(alloc.values())) > 0:
        room = {s: n for s, n in counts.items() if alloc[s] < n}
        if not room:
            break
        share = {s: left * n / sum(room.values()) for s, n in room.items()}
        add = {s: min(int(v), counts[s] - alloc[s]) for s, v in share.items()}
        by_remainder = sorted(room, key=lambda s: (-(share[s] - int(share[s])), -counts[s], s))
        for s in by_remainder[: left - sum(add.values())]:
            add[s] = min(add[s] + 1, counts[s] - alloc[s])
        for s, n in add.items():
            alloc[s] += n
    return alloc


def select_national(plants: dict[int, Plant], exclude_states: Collection[str] = ("FL",),
                    total: int = NATIONAL_SITES) -> list[Plant]:
    """`total` large plants outside `exclude_states`, spread over every state (see allocate()). In each state the
    largest plants come first, skipping any within MIN_SPACING_MI of one already picked."""
    large: dict[str, list[Plant]] = {}
    for p in plants.values():
        if p.capacity_mw >= LARGE_PLANT_MW and p.state not in exclude_states:
            large.setdefault(p.state, []).append(p)
    per_state = allocate({s: len(ps) for s, ps in large.items()}, total)

    chosen: list[Plant] = []
    for state in sorted(large):
        picked: list[Plant] = []
        for p in sorted(large[state], key=lambda p: (-p.capacity_mw, p.id)):
            if len(picked) == per_state[state]:
                break
            if all(miles(p.latitude, p.longitude, q.latitude, q.longitude) >= MIN_SPACING_MI for q in picked):
                picked.append(p)
        chosen.extend(picked)
    return chosen


UPSERT_SQL = """
INSERT INTO sites (id, name, utility, county, state, plant_type, capacity_mw, latitude, longitude, geom, timezone)
VALUES (%(id)s, %(name)s, %(utility)s, %(county)s, %(state)s, %(plant_type)s, %(capacity_mw)s, %(latitude)s,
        %(longitude)s, ST_SetSRID(ST_MakePoint(%(longitude)s, %(latitude)s), 4326)::geography, %(timezone)s)
ON CONFLICT (id) DO UPDATE SET
    name = EXCLUDED.name, utility = EXCLUDED.utility, county = EXCLUDED.county, state = EXCLUDED.state,
    plant_type = EXCLUDED.plant_type, capacity_mw = EXCLUDED.capacity_mw,
    latitude = EXCLUDED.latitude, longitude = EXCLUDED.longitude, geom = EXCLUDED.geom,
    timezone = EXCLUDED.timezone
"""


def write(plants: list[Plant], prune: bool = False) -> list[str]:
    """Upsert the plants. With prune, also delete other sites and return their names.

    Pruning fails (and nothing is deleted) if weather, scores, alerts etc. already reference a site.
    """
    from planit.db import connection

    params = [{**p.__dict__, "timezone": p.timezone} for p in plants]
    with connection() as conn, conn.transaction(), conn.cursor() as cur:
        cur.executemany(UPSERT_SQL, params)
        if not prune:
            return []
        cur.execute("DELETE FROM sites WHERE NOT (id = ANY(%s)) RETURNING name", ([p.id for p in plants],))
        return [r["name"] for r in cur.fetchall()]


def newest_raw_file() -> Path:
    files = sorted(RAW_DIR.glob("*generator*.xlsx"), key=lambda p: p.stat().st_mtime)
    if not files:
        sys.exit(f"No EIA-860M file in {RAW_DIR}. Download one from https://www.eia.gov/electricity/data/eia860m/")
    return files[-1]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("file", nargs="?", type=Path, help="EIA-860M .xlsx (default: newest in data/raw/)")
    parser.add_argument("--write", action="store_true", help="Upsert into the database (default: dry run)")
    parser.add_argument("--prune", action="store_true", help="With --write: delete sites not in the selection")
    args = parser.parse_args()

    path = args.file or newest_raw_file()
    print(f"Reading {path.name} ...")
    every_plant = read_plants(path)
    florida, missing = select(every_plant)
    national = select_national(every_plant)
    plants = florida + national
    for i, p in enumerate(plants, 1):
        note = SELECTED_SITES.get(p.id, "")
        print(f"{i:>3}. {p.id:>6} {p.state} {p.name[:30]:30} {p.plant_type[:28]:28} {p.capacity_mw:>7.1f} MW  "
              f"{p.county[:12]:12} {p.timezone:20} {note}")
    states = Counter(p.state for p in plants)
    print(f"\n{len(plants)} sites ({len(florida)} Florida + {len(national)} national) in {len(states)} states")
    print("Per state: " + " ".join(f"{s}:{n}" for s, n in states.most_common()))
    print("Timezones: " + ", ".join(f"{tz} {n}" for tz, n in Counter(p.timezone for p in plants).most_common()))
    if missing:
        print(f"WARNING: selected plant ids not in this file (retired or renumbered?): {missing}")

    if args.write:
        removed = write(plants, prune=args.prune)
        print(f"Upserted {len(plants)} sites.")
        if removed:
            print(f"Pruned {len(removed)} sites not in the selection: {', '.join(removed)}")
    else:
        print("Dry run: nothing written. Re-run with --write to load into the database.")


if __name__ == "__main__":
    main()
